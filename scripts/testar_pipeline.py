"""
Testa o PIPELINE INTEIRO, de ponta a ponta, com dados INVENTADOS.

Roda, em ordem: gerar dados -> limpeza -> duplicatas -> descritiva -> embeddings
-> BERTopic -> avaliacao -> grid -> planilha de anotacao -> (anotacao SIMULADA)
-> kappa -> rotulos -> figuras. Depois confere se os arquivos esperados existem.

SEGURANCA: so roda com config.USAR_MOCK = True (dados inventados). Ele apaga e
recria arquivos em outputs/mock e data/mock, NUNCA em outputs/real.

Uso (na raiz do projeto, com o ambiente virtual ativo):
    python scripts/testar_pipeline.py                # rapido: vetores de TESTE (TF-IDF), ~1-2 min
    python scripts/testar_pipeline.py --modelo-real  # usa o modelo de embeddings de verdade
                                                     # (baixa ~1 GB na 1a vez; testa o download)
    python scripts/testar_pipeline.py --grid-rapido  # grid minimo (mais rapido ainda)
"""
from __future__ import annotations

import argparse
import subprocess
import sys
import time
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

from src import config  # noqa: E402

VERDE, VERMELHO, RESET = "\033[92m", "\033[91m", "\033[0m"


# ---------------------------------------------------------------------------
# Anotacao SIMULADA (so para testar o codigo; usa a coluna inventada da base mock)
# ---------------------------------------------------------------------------
def simular_anotacao(usar_tfidf: bool) -> None:
    from openpyxl import load_workbook

    from src.anotacao import COL_TATICAS, EIXO1, caminho_anotacao, carregar_resultados_topicos

    assert config.USAR_MOCK, "A simulacao so pode rodar com dados inventados."
    cats = list(EIXO1)
    por_narrativa = {
        "preco_pix": cats[1], "sem_receita": cats[1], "whatsapp": cats[1],
        "resultado_rapido": cats[1], "falso_frete": cats[2], "informativo": cats[0], "ruido": cats[4],
    }
    taticas_narrativa = {
        "preco_pix": ["preco_desconto"], "sem_receita": ["falsa_legitimidade"],
        "whatsapp": ["canal_conversao"], "resultado_rapido": ["resultado_rapido", "depoimento"],
        "falso_frete": ["escassez", "preco_desconto"], "informativo": ["autoridade"], "ruido": [],
    }
    anuncios, _ = carregar_resultados_topicos(usar_tfidf)
    verdade = dict(zip(anuncios["id"], anuncios["narrativa_sintetica"]))
    rng = np.random.default_rng(7)

    def com_erro(valor: str, p: float) -> str:
        return valor if rng.random() > p else str(rng.choice(cats))

    arq = caminho_anotacao(usar_tfidf)
    wb = load_workbook(arq)

    wa = wb["Amostra_topicos"]                       # anotador classifica cada anuncio (5% de erro)
    for r in range(2, wa.max_row + 1):
        wa.cell(r, 6).value = com_erro(por_narrativa[verdade[wa.cell(r, 3).value]], 0.05)
    amostra = pd.DataFrame([[wa.cell(r, 1).value, wa.cell(r, 3).value, wa.cell(r, 6).value]
                            for r in range(2, wa.max_row + 1)], columns=["topico", "id", "cat"])
    maioria = amostra.groupby("topico")["cat"].agg(lambda s: s.value_counts().idxmax()).to_dict()
    narr_do_topico = (anuncios.assign(n=anuncios["narrativa_sintetica"])
                      .groupby("topico")["n"].agg(lambda s: s.value_counts().idxmax()).to_dict())

    wt = wb["Topicos"]
    cabecalho = [c.value for c in wt[1]]
    for r in range(2, wt.max_row + 1):
        t = wt.cell(r, 1).value
        wt.cell(r, cabecalho.index("nome_topico") + 1).value = f"Topico {t} (simulado)"
        wt.cell(r, cabecalho.index("categoria_eixo1") + 1).value = maioria.get(t, cats[4])
        wt.cell(r, cabecalho.index("fraude") + 1).value = "nao" if maioria.get(t) == cats[0] else "sim"
        for tat in taticas_narrativa[narr_do_topico[t]]:
            wt.cell(r, cabecalho.index(f"tatica_{tat}") + 1).value = "x"

    wv = wb["Validacao_cega"]                        # dois anotadores com erros diferentes
    for r in range(2, wv.max_row + 1):
        v = por_narrativa[verdade[wv.cell(r, 1).value]]
        wv.cell(r, 3).value = com_erro(v, 0.05)
        wv.cell(r, 4).value = com_erro(v, 0.12)
    wb.save(arq)
    assert all(c in cabecalho for c in COL_TATICAS)


# ---------------------------------------------------------------------------
# Execucao
# ---------------------------------------------------------------------------
def rodar(args_cmd: list[str]) -> tuple[bool, str]:
    proc = subprocess.run([sys.executable, *args_cmd], cwd=RAIZ, capture_output=True,
                          text=True, encoding="utf-8", errors="replace")
    saida = (proc.stdout or "") + (proc.stderr or "")
    ok = proc.returncode == 0 and "Traceback" not in saida
    return ok, saida


def main() -> int:
    ap = argparse.ArgumentParser(description="Testa o pipeline completo com dados inventados.")
    ap.add_argument("--modelo-real", action="store_true",
                    help="usa o modelo de embeddings de verdade (baixa ~1 GB na 1a vez)")
    ap.add_argument("--grid-rapido", action="store_true", help="grid minimo")
    args = ap.parse_args()

    if not config.USAR_MOCK:
        print(f"{VERMELHO}PARADO:{RESET} config.USAR_MOCK esta False (base REAL). "
              "Este teste so roda com dados inventados. Volte para USAR_MOCK = True.")
        return 2

    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print(f"{VERMELHO}FALTA UM PACOTE:{RESET} instale com  pip install openpyxl  e rode de novo.")
        return 2

    tfidf = not args.modelo_real
    flag_tfidf = ["--tfidf"] if tfidf else []
    prefixo = "bertopic_tfidf_teste" if tfidf else "bertopic"
    sufixo = "_tfidf_teste" if tfidf else ""
    out = config.DIR_SAIDA
    anot = out / ("anotacao_tfidf_teste.xlsx" if tfidf else "anotacao.xlsx")

    # limpa so o que e gerado por este teste (a pasta de saida do modo mock)
    for arq in (anot, out / f"grid_resultados{sufixo}.csv", out / f"grid_escolhido{sufixo}.json"):
        if arq.exists():
            arq.unlink()

    passos: list[tuple[str, list[str], list[Path]]] = [
        ("1. gerar dados inventados", ["scripts/gerar_dados_teste.py"], [config.DIR_MOCK / "anuncios_teste.csv"]),
        ("2. auditoria e limpeza", ["-m", "src.dados", "--forcar"], [config.ARQ_LIMPO]),
        ("3. duplicatas", ["-m", "src.duplicatas"], [out / "anuncios_dedup.csv", out / "anuncios_com_grupos.csv"]),
        ("4. estatistica descritiva", ["-m", "src.descritiva", "--dedup"],
         [out / "descritiva_dedup" / "desc_resumo.csv", out / "descritiva_dedup" / "fig_palavras_frequentes.png"]),
        ("5. embeddings", ["-m", "src.embeddings", "--dedup", *flag_tfidf],
         [out / ("embeddings_tfidf_teste.npy" if tfidf else "embeddings.npy")]),
        ("6. BERTopic", ["-m", "src.modelagem", "--dedup", *flag_tfidf],
         [out / f"{prefixo}_anuncios_com_topicos.csv", out / f"{prefixo}_resumo_topicos.csv"]),
        ("7. avaliacao + estabilidade", ["-m", "src.avaliacao", "--dedup", "--estabilidade", *flag_tfidf], []),
        ("8. grid de hiperparametros",
         ["-m", "src.grid", "--dedup", *flag_tfidf, *(["--rapido"] if args.grid_rapido else [])],
         [out / f"grid_resultados{sufixo}.csv", out / f"grid_escolhido{sufixo}.json"]),
        ("9. planilha de anotacao", ["-m", "src.anotacao", "gerar", *flag_tfidf], [anot]),
        ("10. anotacao SIMULADA", None, []),
        ("11. kappa e pureza", ["-m", "src.anotacao", "concordancia", *flag_tfidf], []),
        ("12. rotulos nos anuncios", ["-m", "src.anotacao", "aplicar", *flag_tfidf],
         [out / f"{prefixo}_anuncios_rotulados.csv"]),
        ("13. figuras", ["-m", "src.figuras", *flag_tfidf],
         [out / "figuras" / n for n in ("grid_hiperparametros.png", "tamanho_topicos.png", "mapa_topicos.png",
                                         "prevalencia_categorias.png", "taticas.png", "categoria_x_tatica.png")]),
    ]

    print(f"TESTE DO PIPELINE COMPLETO  |  embeddings: {'TESTE (TF-IDF)' if tfidf else 'MODELO REAL'}\n")
    inicio, falhas = time.time(), []
    for nome, cmd, esperados in passos:
        t0 = time.time()
        if cmd is None:
            try:
                simular_anotacao(tfidf)
                ok, saida = True, ""
            except Exception as erro:       # noqa: BLE001
                ok, saida = False, f"{type(erro).__name__}: {erro}"
        else:
            ok, saida = rodar(cmd)
        faltando = [p for p in esperados if not p.exists()]
        if faltando:
            ok = False
            saida += "\nArquivos esperados que NAO foram criados:\n" + "\n".join(f"  {p}" for p in faltando)
        print(f"  [{VERDE + ' OK ' + RESET if ok else VERMELHO + 'FALHA' + RESET}] {nome:34s} {time.time() - t0:6.1f}s")
        if not ok:
            falhas.append(nome)
            print("\n".join("      " + l for l in saida.strip().splitlines()[-14:]))

    print(f"\nTempo total: {time.time() - inicio:.0f}s")
    if falhas:
        print(f"{VERMELHO}{len(falhas)} passo(s) falharam:{RESET} {', '.join(falhas)}")
        print("Cole a mensagem de erro acima para pedir ajuda.")
        return 1
    print(f"{VERDE}TUDO CERTO: o pipeline completo funciona.{RESET}")
    print("Lembrete: nada de outputs/mock vale como resultado. Para a base real: USAR_MOCK = False.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
