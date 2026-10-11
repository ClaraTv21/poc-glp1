"""
Anotacao manual: da lista de topicos para a taxonomia da fraude.

O BERTopic agrupa anuncios parecidos, mas QUEM da nome e classifica cada grupo
e voce. Este modulo prepara e confere esse trabalho.

    python -m src.anotacao gerar         cria a planilha de anotacao (.xlsx)
    python -m src.anotacao concordancia  calcula kappa e a pureza dos topicos
    python -m src.anotacao aplicar       leva os rotulos dos topicos para cada anuncio

A planilha tem 5 abas visiveis:
    LEIA-ME           passo a passo
    Codebook          as categorias (Eixo 1) e taticas (Eixo 2), com definicoes
                      em RASCUNHO v0: edite antes de anotar!
    Topicos           1 linha por topico: voce da nome, categoria e taticas
    Amostra_topicos   anuncios sorteados de cada topico: voce classifica um a um
    Validacao_cega    anuncios sorteados SEM mostrar o topico (para o 2o anotador)
(e uma aba oculta com o gabarito da validacao cega).

Eixo 1 = tipo de risco (UMA categoria por anuncio/topico).
Eixo 2 = taticas de persuasao (varias podem estar presentes).

Privacidade: a planilha contem o texto ORIGINAL dos anuncios (pode ter telefones
e nomes). Nao compartilhe fora do grupo de pesquisa e nao cole trechos no relatorio
sem anonimizar. Nao ha nome de anunciante na planilha.
"""
from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
import pandas as pd

try:                                     # planilhas: so e exigido para gerar/ler o .xlsx
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    _TEM_OPENPYXL = True
except ImportError:                      # figuras.py continua funcionando sem o openpyxl
    _TEM_OPENPYXL = False

from src import config
from src.cli import executar  # noqa: E402
from src.modelagem import prefixo_resultados

# ---------------------------------------------------------------------------
# Codebook v0 (RASCUNHO: ajuste as definicoes antes de comecar a anotar)
# ---------------------------------------------------------------------------
SEM_ROTULO = "Sem rótulo"

# categoria -> (definicao, inclui, nao inclui)
EIXO1: dict[str, tuple[str, str, str]] = {
    "Legítimo / informativo": (
        "Anúncio de serviço de saúde regular ou conteúdo educativo, sem sinais de fraude.",
        "Clínica ou farmácia identificada; orientação 'consulte seu médico'; conteúdo educativo; notícia.",
        "Oferta sem receita, preço muito abaixo do mercado, contato só por WhatsApp."),
    "Venda irregular / possível falsificação": (
        "Oferta de medicamento controlado fora do canal regular (sem receita, de origem duvidosa).",
        "'Sem receita'; 'importado'; 'original' com preço absurdo; entrega 'discreta'; caneta manipulada.",
        "Cobrança de taxa falsa para liberar entrega (isso é Estelionato financeiro)."),
    "Estelionato financeiro": (
        "Golpe para tirar dinheiro, em geral sem entregar o produto.",
        "Falso frete ou taxa de liberação; Pix adiantado; site falso imitando órgão oficial; phishing; pedido de dados.",
        "Venda que entrega produto irregular (isso é Venda irregular)."),
    "Exercício ilegal da medicina": (
        "Prescrição, diagnóstico, 'protocolo' ou aplicação feitos por quem não é médico.",
        "'Consulta' sem médico; dose indicada pelo vendedor; 'aplicamos para você'.",
        "Apenas vender o produto (isso é Venda irregular)."),
    "Outro / incerto": (
        "Não cabe nas anteriores ou faltam informações para decidir.",
        "Texto cortado, ambíguo ou fora do tema.",
        "Use com moderação e explique o motivo em 'observacao'."),
}

# chave -> (nome, definicao, exemplo)
EIXO2: dict[str, tuple[str, str, str]] = {
    "autoridade": ("Autoridade", "Cita médico, selo, registro, 'aprovado' ou órgão oficial para dar credibilidade.",
                   "'Aprovado pela Anvisa', 'indicado por endocrinologista'"),
    "resultado_rapido": ("Resultado rápido ou garantido", "Promete emagrecimento rápido, garantido ou sem esforço.",
                         "'Perca 10 kg em 30 dias, resultado garantido'"),
    "depoimento": ("Depoimento / antes e depois", "Usa relatos de clientes ou fotos de antes e depois.",
                   "'Olha como a Maria ficou'"),
    "escassez": ("Escassez ou urgência", "Pressiona com estoque limitado ou prazo curto.",
                 "'Últimas unidades', 'só hoje'"),
    "preco_desconto": ("Preço e desconto", "Destaca preço baixo, promoção, frete grátis ou parcelamento.",
                       "'De R$ 1.200 por R$ 299 no Pix'"),
    "falsa_legitimidade": ("Falsa legitimidade", "Laudos, fotos de laboratório, 'lacrado', 'original' ou notas sem comprovação.",
                           "'Produto original lacrado com nota'"),
    "ofuscacao": ("Ofuscação textual", "Escrita alterada para fugir da moderação.",
                  "'O z e m p i c', '0zempic', 'Mounj@ro'"),
    "canal_conversao": ("Canal de conversão", "Leva a conversa para fora da plataforma.",
                        "'Chame no WhatsApp', 'link na bio', 'Telegram'"),
}

DESEMPATE = ("Regra de desempate (rascunho): se o anúncio couber em mais de uma categoria, use a mais grave, "
             "nesta ordem: Estelionato financeiro > Exercício ilegal da medicina > Venda irregular > "
             "Legítimo. Explique em 'observacao'.")

_LANDIS_KOCH = [(0.0, "pobre"), (0.20, "leve"), (0.40, "razoável"), (0.60, "moderada"),
                (0.80, "substancial"), (1.01, "quase perfeita")]

COL_TATICAS = [f"tatica_{k}" for k in EIXO2]


# ---------------------------------------------------------------------------
# Caminhos
# ---------------------------------------------------------------------------
def caminho_anotacao(usar_tfidf: bool = False) -> Path:
    return config.DIR_SAIDA / ("anotacao_tfidf_teste.xlsx" if usar_tfidf else "anotacao.xlsx")


def caminho_rotulados(usar_tfidf: bool = False) -> Path:
    return config.DIR_SAIDA / f"{prefixo_resultados(usar_tfidf)}_anuncios_rotulados.csv"


def carregar_resultados_topicos(usar_tfidf: bool = False) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Le os CSVs salvos por modelagem.salvar_resultados: (anuncios com topico, resumo dos topicos)."""
    prefixo = prefixo_resultados(usar_tfidf)
    arq_anuncios = config.DIR_SAIDA / f"{prefixo}_anuncios_com_topicos.csv"
    arq_resumo = config.DIR_SAIDA / f"{prefixo}_resumo_topicos.csv"
    if not arq_anuncios.exists():
        raise FileNotFoundError(
            f"Nao encontrei {arq_anuncios}.\nRode antes o modelo: python -m src.modelagem --dedup"
            + (" --tfidf" if usar_tfidf else ""))
    anuncios = pd.read_csv(arq_anuncios, encoding="utf-8")
    anuncios["texto"] = anuncios["texto"].fillna("").astype(str)
    resumo = (pd.read_csv(arq_resumo, encoding="utf-8-sig") if arq_resumo.exists()
              else anuncios.groupby("topico").size().rename("n_anuncios").reset_index())
    return anuncios, resumo


# ---------------------------------------------------------------------------
# Geracao da planilha
# ---------------------------------------------------------------------------
_CABECALHO = PatternFill("solid", fgColor="D9E2F3") if _TEM_OPENPYXL else None
_PREENCHER = PatternFill("solid", fgColor="FFF2CC") if _TEM_OPENPYXL else None   # amarelo = preencher


def _exigir_openpyxl() -> None:
    if not _TEM_OPENPYXL:
        raise ImportError("Para criar/ler a planilha de anotacao instale: pip install openpyxl")


def _estilo_cabecalho(ws, n_colunas: int, colunas_preencher: set[int] | None = None) -> None:
    for c in range(1, n_colunas + 1):
        cel = ws.cell(1, c)
        cel.font = Font(name="Arial", bold=True, size=10)
        cel.fill = _PREENCHER if colunas_preencher and c in colunas_preencher else _CABECALHO
        cel.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"


def _estilo_corpo(ws, larguras: list[int]) -> None:
    for i, w in enumerate(larguras, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for linha in ws.iter_rows(min_row=2):
        for cel in linha:
            cel.font = Font(name="Arial", size=10)
            cel.alignment = Alignment(wrap_text=True, vertical="top")


def _lista(ws, formula: str, intervalo: str, aviso: str) -> None:
    dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                        showErrorMessage=True, errorTitle="Valor invalido", error=aviso)
    ws.add_data_validation(dv)
    dv.add(intervalo)


def gerar_planilha(anuncios: pd.DataFrame, resumo: pd.DataFrame, caminho: str | Path,
                   n_por_topico: int = 15, n_validacao: int = 100, n_outliers: int = 10,
                   seed: int = config.SEED) -> Path:
    """Cria o .xlsx de anotacao. `anuncios` precisa das colunas id, texto e topico."""
    _exigir_openpyxl()
    for col in ("id", "texto", "topico"):
        if col not in anuncios.columns:
            raise ValueError(f"A coluna '{col}' nao existe nos anuncios com topico.")
    anuncios = anuncios.reset_index(drop=True)
    rng = np.random.default_rng(seed)

    # amostra por topico
    partes = []
    for topico, sub in anuncios.groupby("topico", sort=True):
        n = n_outliers if topico == -1 else n_por_topico
        n = min(n, len(sub))
        idx = rng.choice(sub.index.to_numpy(), size=n, replace=False)
        partes.append(anuncios.loc[sorted(idx)])
    amostra = pd.concat(partes) if partes else anuncios.head(0)

    # amostra cega: sorteada entre os que NAO estao na amostra por topico
    restantes = anuncios.drop(index=amostra.index)
    n_val = min(n_validacao, len(restantes))
    if n_val < n_validacao:
        print(f"Aviso: so {n_val} anuncios sobraram para a validacao cega (voce pediu {n_validacao}). "
              "Reduza --n-por-topico para liberar mais anuncios.")
    cega = anuncios.loc[sorted(rng.choice(restantes.index.to_numpy(), size=n_val, replace=False))] \
        if n_val else restantes.head(0)

    wb = Workbook()

    # --- LEIA-ME
    ws = wb.active
    ws.title = "LEIA-ME"
    linhas = [
        "ANOTAÇÃO DOS ANÚNCIOS — passo a passo",
        "",
        "0. Leia a aba Codebook e AJUSTE as definições ao seu critério ANTES de anotar (elas são rascunho).",
        "1. Aba Topicos: para cada tópico, leia as palavras-chave, o exemplo e a amostra na aba Amostra_topicos.",
        "2. Aba Amostra_topicos: classifique CADA anúncio em 'categoria_eixo1' (lista suspensa). Se estiver em dúvida, marque 'incerto'.",
        "3. Volte à aba Topicos e preencha: nome_topico, descricao, categoria_eixo1 (a que predomina), fraude (sim/nao/incerto) e as táticas presentes (x).",
        "4. Aba Validacao_cega: classifique os anúncios SEM olhar os tópicos. Um 2º anotador deve preencher 'categoria_anotador2' numa CÓPIA do arquivo (mesma ordem de linhas).",
        "5. Rode: python -m src.anotacao concordancia   (calcula kappa e a pureza dos tópicos)",
        "6. Rode: python -m src.anotacao aplicar        (leva os rótulos dos tópicos para cada anúncio)",
        "",
        "Células AMARELAS no cabeçalho = colunas para você preencher.",
        "Não clique em links nem teste telefones dos anúncios. Não pesquise o anúncio na internet durante a validação cega.",
        "Esta planilha contém textos originais (podem ter telefones): não compartilhe fora do grupo de pesquisa.",
        "Salve com Ctrl+S ao terminar e NÃO mude a ordem das linhas nem apague colunas.",
    ]
    if config.USAR_MOCK:
        linhas.insert(1, ">>> DADOS INVENTADOS (teste do código). Não use nada daqui no relatório. <<<")
    for i, texto in enumerate(linhas, 1):
        ws.cell(i, 1, texto).font = Font(name="Arial", size=10, bold=(i == 1))
    ws.column_dimensions["A"].width = 150

    # --- Codebook
    cb = wb.create_sheet("Codebook")
    cb["A1"] = "CODEBOOK v0 (rascunho) — edite as definições antes de anotar"
    cb["A1"].font = Font(name="Arial", bold=True)
    for c, t in enumerate(["Categoria (Eixo 1: tipo de risco)", "Definição", "Inclui", "Não inclui"], 1):
        cel = cb.cell(3, c, t)
        cel.font, cel.fill = Font(name="Arial", bold=True, size=10), _CABECALHO
    for i, (cat, (definicao, inclui, nao)) in enumerate(EIXO1.items()):
        for c, t in enumerate([cat, definicao, inclui, nao], 1):
            cb.cell(4 + i, c, t)
    fim1 = 3 + len(EIXO1)
    cb.cell(fim1 + 2, 1, DESEMPATE)
    ini2 = fim1 + 4
    for c, t in enumerate(["Tática (Eixo 2: persuasão)", "Definição", "Exemplo"], 1):
        cel = cb.cell(ini2, c, t)
        cel.font, cel.fill = Font(name="Arial", bold=True, size=10), _CABECALHO
    for i, (_, (nome, definicao, exemplo)) in enumerate(EIXO2.items()):
        for c, t in enumerate([nome, definicao, exemplo], 1):
            cb.cell(ini2 + 1 + i, c, t)
    for linha in cb.iter_rows(min_row=3):
        for cel in linha:
            if cel.font.name != "Arial" or not cel.font.bold:
                cel.font = Font(name="Arial", size=10, bold=cel.font.bold)
            cel.alignment = Alignment(wrap_text=True, vertical="top")
    for letra, w in zip("ABCD", (38, 60, 60, 50)):
        cb.column_dimensions[letra].width = w
    lista_cat = f"=Codebook!$A$4:$A${fim1}"
    msg_cat = "Escolha uma categoria da lista (veja a aba Codebook)."

    # --- Topicos
    wt = wb.create_sheet("Topicos")
    base_cols = ["topico", "n_anuncios", "pct", "palavras_chave", "exemplo"]
    resumo = resumo.copy()
    for c in base_cols:
        if c not in resumo.columns:
            resumo[c] = ""
    preencher = ["nome_topico", "descricao", "categoria_eixo1", "fraude"] + COL_TATICAS + ["observacao"]
    wt.append(base_cols + preencher)
    for _, r in resumo.sort_values("topico").iterrows():
        wt.append([int(r["topico"]), int(r["n_anuncios"]), r["pct"], str(r["palavras_chave"]), str(r["exemplo"])]
                  + [""] * len(preencher))
    _estilo_cabecalho(wt, len(base_cols) + len(preencher),
                      set(range(len(base_cols) + 1, len(base_cols) + len(preencher) + 1)))
    _estilo_corpo(wt, [8, 10, 7, 38, 50, 24, 40, 30, 12] + [11] * len(COL_TATICAS) + [30])
    ultima = wt.max_row
    col_cat = base_cols.__len__() + 3                     # categoria_eixo1
    L = wt.cell(1, col_cat).column_letter
    _lista(wt, lista_cat, f"{L}2:{L}{ultima}", msg_cat)
    L = wt.cell(1, col_cat + 1).column_letter              # fraude
    _lista(wt, '"sim,nao,incerto"', f"{L}2:{L}{ultima}", "Use: sim, nao ou incerto.")
    ini_t = col_cat + 2
    for k in range(len(COL_TATICAS)):
        L = wt.cell(1, ini_t + k).column_letter
        _lista(wt, '"x"', f"{L}2:{L}{ultima}", "Marque apenas com x (ou deixe vazio).")

    # --- Amostra_topicos
    wa = wb.create_sheet("Amostra_topicos")
    cols_a = ["topico", "palavras_do_topico", "id_anuncio", "texto", "n_copias_grupo",
              "categoria_eixo1", "incerto", "observacao"]
    wa.append(cols_a)
    palavras = dict(zip(resumo["topico"].astype(int), resumo["palavras_chave"].astype(str)))
    for _, r in amostra.iterrows():
        wa.append([int(r["topico"]), palavras.get(int(r["topico"]), ""), r["id"], r["texto"],
                   int(r["n_copias_grupo"]) if "n_copias_grupo" in r and pd.notna(r.get("n_copias_grupo")) else 1,
                   "", "", ""])
    _estilo_cabecalho(wa, len(cols_a), {6, 7, 8})
    _estilo_corpo(wa, [8, 32, 11, 80, 10, 32, 9, 30])
    n = max(wa.max_row, 2)
    _lista(wa, lista_cat, f"F2:F{n}", msg_cat)
    _lista(wa, '"x"', f"G2:G{n}", "Marque apenas com x (ou deixe vazio).")

    # --- Validacao_cega (sem topico!)
    wv = wb.create_sheet("Validacao_cega")
    cols_v = ["id_anuncio", "texto", "categoria_anotador1", "categoria_anotador2", "incerto", "observacao"]
    wv.append(cols_v)
    for _, r in cega.iterrows():
        wv.append([r["id"], r["texto"], "", "", "", ""])
    _estilo_cabecalho(wv, len(cols_v), {3, 4, 5, 6})
    _estilo_corpo(wv, [11, 90, 32, 32, 9, 30])
    n = max(wv.max_row, 2)
    _lista(wv, lista_cat, f"C2:D{n}", msg_cat)
    _lista(wv, '"x"', f"E2:E{n}", "Marque apenas com x (ou deixe vazio).")

    # --- gabarito oculto (topico de cada anuncio da validacao cega)
    wg = wb.create_sheet("_gabarito")
    wg.append(["id_anuncio", "topico"])
    for _, r in cega.iterrows():
        wg.append([r["id"], int(r["topico"])])
    wg.sheet_state = "hidden"

    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)
    return caminho


# ---------------------------------------------------------------------------
# Leitura da planilha preenchida
# ---------------------------------------------------------------------------
def _marcado(serie: pd.Series) -> pd.Series:
    return serie.astype(str).str.strip().str.lower().isin(["x", "sim", "s", "1", "true"])


def _texto_ou_nan(serie: pd.Series) -> pd.Series:
    s = serie.astype("string").str.strip()
    return s.mask(s.isin(["", "nan", "<NA>"]))


def ler_rotulos_topicos(caminho: str | Path) -> pd.DataFrame:
    """Aba Topicos -> tabela: topico, nome_topico, descricao, categoria_eixo1, fraude, tatica_*."""
    _exigir_openpyxl()
    t = pd.read_excel(caminho, sheet_name="Topicos", engine="openpyxl")
    t["topico"] = t["topico"].astype(int)
    for col in ("nome_topico", "descricao", "categoria_eixo1", "fraude", "observacao"):
        t[col] = _texto_ou_nan(t[col]) if col in t.columns else pd.NA
    for col in COL_TATICAS:
        t[col] = _marcado(t[col]) if col in t.columns else False
    return t[["topico", "nome_topico", "descricao", "categoria_eixo1", "fraude", *COL_TATICAS]]


def rotular_anuncios(anuncios: pd.DataFrame, rotulos: pd.DataFrame) -> pd.DataFrame:
    """Copia para cada anuncio o rotulo do seu topico (categoria, fraude, taticas)."""
    r = rotulos.copy()
    r["categoria_eixo1"] = r["categoria_eixo1"].astype("object").where(r["categoria_eixo1"].notna(), SEM_ROTULO)
    r = r.rename(columns={"categoria_eixo1": "categoria"})
    saida = anuncios.merge(r, on="topico", how="left")
    saida["categoria"] = saida["categoria"].fillna(SEM_ROTULO)
    for col in COL_TATICAS:
        saida[col] = saida[col].fillna(False).astype(bool)
    return saida


def _interpretar_kappa(k: float) -> str:
    if k != k:
        return "indefinido"
    for limite, nome in _LANDIS_KOCH:
        if k < limite:
            return nome
    return "quase perfeita"


def _kappa(a: pd.Series, b: pd.Series) -> dict:
    from sklearn.metrics import cohen_kappa_score
    ok = a.notna() & b.notna()
    n = int(ok.sum())
    if n < 2:
        return {"n": n, "concordancia_simples": float("nan"), "kappa": float("nan"), "leitura": "indefinido"}
    a, b = a[ok].astype(str), b[ok].astype(str)
    k = float(cohen_kappa_score(a, b)) if (a.nunique() > 1 or b.nunique() > 1) else float("nan")
    return {"n": n, "concordancia_simples": round(float((a == b).mean()), 3),
            "kappa": round(k, 3) if k == k else float("nan"), "leitura": _interpretar_kappa(k)}


def pureza_dos_topicos(caminho: str | Path) -> pd.DataFrame:
    """
    Dos anuncios que VOCE classificou na amostra, que % tem a mesma categoria
    do topico? (topico homogeneo = pureza alta). Topicos sem rotulo ficam de fora.
    """
    amostra = pd.read_excel(caminho, sheet_name="Amostra_topicos", engine="openpyxl")
    rot = ler_rotulos_topicos(caminho)[["topico", "categoria_eixo1"]].rename(
        columns={"categoria_eixo1": "categoria_topico"})
    amostra["categoria_anuncio"] = _texto_ou_nan(amostra["categoria_eixo1"])
    amostra = amostra.merge(rot, on="topico", how="left")
    amostra = amostra[amostra["categoria_anuncio"].notna() & amostra["categoria_topico"].notna()]
    if amostra.empty:
        return pd.DataFrame(columns=["topico", "n_classificados", "pureza_pct", "categoria_topico"])
    saida = (amostra.assign(igual=amostra["categoria_anuncio"] == amostra["categoria_topico"])
             .groupby(["topico", "categoria_topico"], as_index=False)
             .agg(n_classificados=("igual", "size"), pureza_pct=("igual", lambda s: round(100 * s.mean(), 1))))
    return saida.sort_values("pureza_pct").reset_index(drop=True)


def calcular_concordancia(caminho: str | Path) -> dict:
    """
    Resultados de validacao:
      anotador1_vs_anotador2   concordancia entre duas pessoas (kappa de Cohen)
      anotador1_vs_topicos     o rotulo que voce deu na validacao cega bate com o
                               rotulo do TOPICO do anuncio? (valida a taxonomia)
    """
    v = pd.read_excel(caminho, sheet_name="Validacao_cega", engine="openpyxl")
    a1, a2 = _texto_ou_nan(v["categoria_anotador1"]), _texto_ou_nan(v["categoria_anotador2"])
    resultado = {"anotador1_vs_anotador2": _kappa(a1, a2)}

    gab = pd.read_excel(caminho, sheet_name="_gabarito", engine="openpyxl")
    rot = ler_rotulos_topicos(caminho)[["topico", "categoria_eixo1"]]
    v = v.assign(categoria_anotador1=a1).merge(gab, on="id_anuncio", how="left").merge(rot, on="topico", how="left")
    resultado["anotador1_vs_topicos"] = _kappa(v["categoria_anotador1"], _texto_ou_nan(v["categoria_eixo1"]))
    return resultado


# ---------------------------------------------------------------------------
# Terminal
# ---------------------------------------------------------------------------
def main() -> None:
    ap = argparse.ArgumentParser(description="Planilha de anotacao, kappa e rotulos.")
    sub = ap.add_subparsers(dest="comando", required=True)
    for nome in ("gerar", "concordancia", "aplicar"):
        sp = sub.add_parser(nome)
        sp.add_argument("--tfidf", action="store_true", help="usa os resultados do teste TF-IDF")
        sp.add_argument("--arquivo", type=Path, default=None, help="caminho do .xlsx")
        if nome == "gerar":
            sp.add_argument("--n-por-topico", type=int, default=15)
            sp.add_argument("--n-validacao", type=int, default=100)
    args = ap.parse_args()
    arquivo = args.arquivo or caminho_anotacao(args.tfidf)
    pd.set_option("display.width", 200)

    if args.comando == "gerar":
        if arquivo.exists():
            raise SystemExit(f"{arquivo} ja existe. Para nao apagar sua anotacao, renomeie ou apague "
                             "o arquivo antes de gerar de novo.")
        anuncios, resumo = carregar_resultados_topicos(args.tfidf)
        gerar_planilha(anuncios, resumo, arquivo, args.n_por_topico, args.n_validacao)
        print(f"Planilha criada: {arquivo}")
        print("Abra no LibreOffice/Excel e siga a aba LEIA-ME.")

    elif args.comando == "concordancia":
        print("PUREZA DOS TOPICOS (anuncios que voce classificou x categoria do topico)")
        pureza = pureza_dos_topicos(arquivo)
        print(pureza.to_string(index=False) if len(pureza) else "  (preencha as abas Topicos e Amostra_topicos)")
        print("\nCONCORDANCIA")
        for nome, r in calcular_concordancia(arquivo).items():
            print(f"  {nome}: {r}")
        print("\nLeitura do kappa (Landis & Koch): <0,20 leve | 0,21-0,40 razoavel | 0,41-0,60 moderada "
              "| 0,61-0,80 substancial | >0,80 quase perfeita")

    elif args.comando == "aplicar":
        anuncios, _ = carregar_resultados_topicos(args.tfidf)
        rotulados = rotular_anuncios(anuncios, ler_rotulos_topicos(arquivo))
        saida = caminho_rotulados(args.tfidf)
        rotulados.to_csv(saida, index=False, encoding="utf-8")
        print(f"Anuncios rotulados salvos em {saida}")
        print(rotulados["categoria"].value_counts().to_string())


if __name__ == "__main__":
    executar(main)
